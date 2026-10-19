"""Tier-0 fast path: DOCX / PPTX / XLSX / HTML → AST directly from the file structure (no OCR, no GPU).

These formats have no page raster, so each page also gets an HTML preview rendered from the original file
(formatting, merged cells, slide layout). Every element's block carries ``data-el`` and every table cell
``data-r``/``data-c`` (AST coordinates) so the verification UI can sync selection with the extracted data.
"""

from __future__ import annotations

import base64
import datetime as dt
import io
from collections.abc import Callable
from html import escape

from idp.assemble.page_builder import table_element, text_markdown
from idp.core.ast import DocStatus, Document, DocumentMetadata, Element, ElementType, Page, TableCell
from idp.ingest.detect import MIME, FileKind
from idp.ingest.legacy_fonts import convert_legacy
from idp.table.builder import finalize_cells, line_maps

Item = tuple[Element, str, str]  # element, preview fragment, extra CSS for its block


def _text_el(type_: ElementType, text: str, level: int | None = None) -> Element:
    text = " ".join(convert_legacy(text).split()) if type_ != ElementType.PARAGRAPH else convert_legacy(text).strip()
    return Element(
        element_id="",
        type=type_,
        reading_order=0,
        text=text,
        level=level,
        content_markdown=text_markdown(type_, text, level),
        extraction_engine="native",
    )


def _table_el(cells: list[TableCell]) -> Element | None:
    cells = finalize_cells(cells, ocr=False)
    if not cells:
        return None
    return table_element(cells, element_id="", reading_order=0, extraction_engine="native")


def _text_html(text: str) -> str:
    return escape(text).replace("\n", "<br>")


def _table_html(
    cells: list[TableCell],
    style_of: Callable[[int, int], str] | None = None,
    col_widths: list[int] | None = None,
    cls: str = "",
) -> str:
    """Original grid (empty lines kept); cells tagged with their coordinates in the finalized AST table."""
    rmap, cmap = line_maps(cells)
    by_row: dict[int, list[TableCell]] = {}
    for c in cells:
        by_row.setdefault(c.row, []).append(c)
    out = [f'<table class="{cls}">' if cls else "<table>"]
    if col_widths:
        out.append("<colgroup>" + "".join(f'<col style="width:{w}px">' for w in col_widths) + "</colgroup>")
    for r in range(max(c.row + c.rowspan for c in cells)):
        out.append("<tr>")
        for c in sorted(by_row.get(r, []), key=lambda c: c.col):
            attrs = (f' rowspan="{c.rowspan}"' if c.rowspan > 1 else "") + (f' colspan="{c.colspan}"' if c.colspan > 1 else "")
            if style := (style_of(c.row, c.col) if style_of else ""):
                attrs += f' style="{style}"'
            ar = next((rmap[i] for i in range(c.row, c.row + c.rowspan) if i in rmap), None)
            ac = next((cmap[k] for k in range(c.col, c.col + c.colspan) if k in cmap), None)
            if ar is not None and ac is not None:
                attrs += f' data-r="{ar}" data-c="{ac}"'
            out.append(f"<td{attrs}>{_text_html(c.text)}</td>")
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


def _table(cells: list[TableCell], **kw) -> tuple[Element | None, str]:
    html = _table_html(cells, **kw) if cells else ""  # before finalize: it renumbers copies
    return _table_el(cells), html


def _page(n: int, items: list[Item], before: str = "", after: str = "") -> tuple[Page, str]:
    elements, blocks = [], []
    for i, (el, frag, style) in enumerate(items):
        el.reading_order = i
        el.element_id = f"el_p{n}_{i + 1:03d}"
        elements.append(el)
        style_attr = f' style="{style}"' if style else ""
        blocks.append(f'<div class="blk t-{el.type.value}" data-el="{el.element_id}"{style_attr}>{frag}</div>')
    return Page(page_number=n, width=0, height=0, elements=elements), before + "".join(blocks) + after


_PREVIEW_CSS = """
body{margin:0;padding:24px;font:15px/1.5 "Times New Roman",serif;color:#111;background:#fff}
.blk{margin:0 0 .5em;border-radius:2px;outline:2px solid transparent;outline-offset:2px;cursor:pointer}
.blk:hover{outline-color:rgba(37,99,235,.35)}
.blk.rev{outline-color:rgba(217,119,6,.55)}
.blk.sel{outline-color:#2563eb;background:rgba(37,99,235,.06)}
h1,h2,h3,h4,h5,h6,p,ul{margin:0}
table{border-collapse:collapse}
td{border:1px solid #9a9a9a;padding:2px 6px;vertical-align:top}
td.cell-hl{outline:3px solid #f59e0b;outline-offset:-2px;background:rgba(245,158,11,.15)}
.sheet-tab{display:inline-block;padding:2px 12px;font:600 13px system-ui,sans-serif;background:#e8f0e8;border:1px solid #9a9a9a;border-radius:4px 4px 0 0}
table.grid{table-layout:fixed;font:13px Calibri,Arial,sans-serif}
table.grid td{border-color:#d0d0d0;white-space:pre-wrap;overflow:hidden}
body.kind-pptx{background:#d9d9d9}
.slide{position:relative;width:100%;container-type:inline-size;overflow:hidden;background:#fff;box-shadow:0 1px 4px rgba(0,0,0,.3);font-family:Calibri,Arial,sans-serif;line-height:1.2}
.slide .blk{position:absolute;margin:0;overflow:hidden}
.slide .pic{position:absolute;object-fit:contain}
.slide td{font-size:inherit}
"""


def _preview_document(body: str, kind: FileKind) -> str:
    return (
        '<!DOCTYPE html><html lang="vi"><head><meta charset="utf-8">'
        f"<style>{_PREVIEW_CSS}</style></head><body class=\"kind-{kind.value}\">{body}</body></html>"
    )


# DOCX -----------------------------------------------------------------------
def _docx_run_font(run, par):
    return run.font.name or (par.style.font.name if par.style is not None and par.style.font is not None else None)


def _docx_paragraph_text(par) -> str:
    parts = []
    for run in par.runs:
        font = _docx_run_font(run, par)
        parts.append(convert_legacy(run.text, font) if font else run.text)
    return "".join(parts) if parts else par.text


def _docx_paragraph_html(par) -> str:
    if not par.runs:
        return _text_html(par.text)
    parts = []
    for run in par.runs:
        font = _docx_run_font(run, par)
        h = _text_html(convert_legacy(run.text, font) if font else run.text)
        if run.bold:
            h = f"<b>{h}</b>"
        if run.italic:
            h = f"<i>{h}</i>"
        if run.underline:
            h = f"<u>{h}</u>"
        parts.append(h)
    return "".join(parts)


_DOCX_ALIGN = {1: "center", 2: "right", 3: "justify"}


def _docx_table(tbl, doc) -> list[TableCell]:
    from docx.table import _Cell

    cells: list[TableCell] = []
    active: dict[int, TableCell] = {}  # grid column → cell open for vertical merge
    for r, tr in enumerate(tbl._tbl.tr_lst):
        col = 0
        for tc in tr.tc_lst:
            span = tc.grid_span
            vmerge = tc.vMerge
            if vmerge == "continue" and col in active:
                active[col].rowspan += 1
            else:
                text = "\n".join(_docx_paragraph_text(p) for p in _Cell(tc, tbl).paragraphs).strip()
                cell = TableCell(row=r, col=col, colspan=span, text=convert_legacy(text))
                cells.append(cell)
                for k in range(col, col + span):
                    active.pop(k, None)
                if vmerge == "restart":
                    active[col] = cell
            col += span
    return cells


def _parse_docx(data: bytes) -> list[tuple[Page, str]]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    items: list[Item] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            par = Paragraph(child, d)
            text = _docx_paragraph_text(par)
            if not text.strip():
                continue
            align = _DOCX_ALIGN.get(int(par.paragraph_format.alignment)) if par.paragraph_format.alignment is not None else None
            inner = _docx_paragraph_html(par)

            def wrap(tag: str, inner: str = inner, align: str | None = align) -> str:
                return f'<{tag} style="text-align:{align}">{inner}</{tag}>' if align else f"<{tag}>{inner}</{tag}>"

            style = (par.style.name if par.style is not None else "").lower()
            if style.startswith("title"):
                items.append((_text_el(ElementType.TITLE, text, 0), wrap("h1"), ""))
            elif style.startswith("heading"):
                digits = "".join(ch for ch in style if ch.isdigit())
                level = int(digits) if digits else 1
                items.append((_text_el(ElementType.HEADING, text, level), wrap(f"h{min(level + 1, 6)}"), ""))
            elif "list" in style or par._p.pPr is not None and par._p.pPr.numPr is not None:
                items.append((_text_el(ElementType.LIST_ITEM, text), f"<ul>{wrap('li')}</ul>", ""))
            else:
                items.append((_text_el(ElementType.PARAGRAPH, text), wrap("p"), ""))
        elif tag == "tbl":
            el, html = _table(_docx_table(Table(child, d), d))
            if el:
                items.append((el, html, ""))
    return [_page(1, items)]


# XLSX -----------------------------------------------------------------------
def format_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if isinstance(v, int):
            body = f"{abs(v):,}".replace(",", ".")
        else:
            body = f"{abs(v):,.10f}".rstrip("0").rstrip(".")
            body = body.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
        return f"-{body}" if v < 0 else body
    return convert_legacy(str(v)).strip()


def _xlsx_color(color) -> str | None:
    rgb = getattr(color, "rgb", None) if color is not None and getattr(color, "type", None) == "rgb" else None
    return f"#{rgb[-6:].upper()}" if isinstance(rgb, str) and len(rgb) >= 6 else None


def _xlsx_style(cell) -> str:
    css = []
    if (f := cell.font) is not None:
        if f.b:
            css.append("font-weight:bold")
        if f.i:
            css.append("font-style:italic")
        if f.u:
            css.append("text-decoration:underline")
        if (c := _xlsx_color(f.color)) and c != "#000000":
            css.append(f"color:{c}")
    if cell.fill is not None and cell.fill.fill_type == "solid" and (c := _xlsx_color(cell.fill.fgColor)):
        css.append(f"background:{c}")
    h = cell.alignment.horizontal if cell.alignment is not None else None
    if h in ("left", "center", "right", "justify"):
        css.append(f"text-align:{h}")
    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
        css.append("text-align:right")
    return ";".join(css)


def _parse_xlsx(data: bytes) -> list[tuple[Page, str]]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    pages: list[tuple[Page, str]] = []
    for ws in wb.worksheets:
        merged = {(m.min_row, m.min_col): (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1) for m in ws.merged_cells.ranges}
        covered = {
            (r, c)
            for m in ws.merged_cells.ranges
            for r in range(m.min_row, m.max_row + 1)
            for c in range(m.min_col, m.max_col + 1)
            if (r, c) != (m.min_row, m.min_col)
        }
        cells = []
        styles: dict[tuple[int, int], str] = {}
        for row in ws.iter_rows(min_row=ws.min_row, max_row=ws.max_row, min_col=ws.min_column, max_col=ws.max_column):
            for cell in row:
                key = (cell.row, cell.column)
                if key in covered:
                    continue
                rs, cs = merged.get(key, (1, 1))
                r, c = cell.row - ws.min_row, cell.column - ws.min_column
                cells.append(TableCell(row=r, col=c, rowspan=rs, colspan=cs, text=format_value(cell.value)))
                styles[(r, c)] = _xlsx_style(cell)
        if not any(c.text for c in cells):
            continue
        widths = []
        for k in range(ws.min_column, ws.max_column + 1):
            letter = get_column_letter(k)
            w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
            widths.append(round((w or 8.43) * 7 + 5))
        grid = _table_html(cells, style_of=lambda r, c: styles.get((r, c), ""), col_widths=widths, cls="grid")
        table = _table_el(cells)
        if table is None:
            continue
        title = (_text_el(ElementType.HEADING, ws.title, 1), f'<div class="sheet-tab">{escape(ws.title)}</div>', "margin:0")
        pages.append(_page(len(pages) + 1, [title, (table, grid, "")]))
    return pages


# PPTX -----------------------------------------------------------------------
def _pptx_text_html(tf, pt_w: float) -> str:
    out = []
    for p in tf.paragraphs:
        runs = []
        for run in p.runs:
            h = _text_html(convert_legacy(run.text, run.font.name) if run.font.name else run.text)
            css = []
            if run.font.bold:
                css.append("font-weight:bold")
            if run.font.italic:
                css.append("font-style:italic")
            if run.font.underline:
                css.append("text-decoration:underline")
            if run.font.size is not None:
                css.append(f"font-size:{run.font.size.pt / pt_w * 100:.3f}cqw")
            try:
                if run.font.color.type is not None and run.font.color.rgb is not None:
                    css.append(f"color:#{run.font.color.rgb}")
            except AttributeError:  # theme colours have no rgb
                pass
            runs.append(f'<span style="{";".join(css)}">{h}</span>' if css else h)
        align = {2: "center", 3: "right", 4: "justify"}.get(int(p.alignment)) if p.alignment is not None else None
        style = f' style="text-align:{align}"' if align else ""
        out.append(f"<p{style}>{''.join(runs) or '&nbsp;'}</p>")
    return "".join(out)


def _parse_pptx(data: bytes) -> list[tuple[Page, str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(io.BytesIO(data))
    sw, sh = prs.slide_width or 9144000, prs.slide_height or 6858000
    pt_w = sw / 12700

    def box(shape) -> str:
        left, top, width, height = (v or 0 for v in (shape.left, shape.top, shape.width, shape.height))
        return f"left:{round(100 * left / sw, 2)}%;top:{round(100 * top / sh, 2)}%;width:{round(100 * width / sw, 2)}%;height:{round(100 * height / sh, 2)}%"

    pages = []
    for n, slide in enumerate(prs.slides, start=1):
        items: list[Item] = []
        pictures = []
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for shape in shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                cells = []
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        if cell.is_spanned:
                            continue
                        rs, cs = (cell.span_height, cell.span_width) if cell.is_merge_origin else (1, 1)
                        cells.append(TableCell(row=r, col=c, rowspan=rs, colspan=cs, text=convert_legacy(cell.text).strip()))
                el, html = _table(cells)
                if el:
                    items.append((el, html, box(shape)))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                is_title = shape.is_placeholder and shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
                el = _text_el(ElementType.TITLE if is_title else ElementType.PARAGRAPH, shape.text_frame.text, 0 if is_title else None)
                style = box(shape) + (f";font-size:{32 / pt_w * 100:.3f}cqw;font-weight:bold" if is_title else "")
                items.append((el, _pptx_text_html(shape.text_frame, pt_w), style))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = shape.image
                except (AttributeError, ValueError, KeyError):  # linked / missing image part
                    continue
                if img.content_type in ("image/png", "image/jpeg", "image/gif"):
                    src = f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"
                    pictures.append(f'<img class="pic" alt="" style="{box(shape)}" src="{src}">')
        before = f'<div class="slide" style="aspect-ratio:{sw}/{sh};font-size:{18 / pt_w * 100:.3f}cqw">' + "".join(pictures)
        pages.append(_page(n, items, before, "</div>"))
    return pages


# HTML -----------------------------------------------------------------------
_BLOCK_CONTAINERS = {"body", "div", "section", "article", "main", "ul", "ol", "header", "footer", "aside", "form", "center", "blockquote"}
_ACTIVE_TAGS = ["iframe", "frame", "frameset", "object", "embed", "applet", "link", "meta", "base"]
_SAFE_ATTRS = {"colspan", "rowspan", "href", "src", "alt", "title", "style", "align", "valign", "width", "height", "border", "bgcolor", "color", "size", "face"}


def _sanitize(soup) -> None:
    """Keep the original markup for the preview, minus anything that can run, load or navigate."""
    for t in soup(_ACTIVE_TAGS):
        t.decompose()
    for tag in soup.find_all(True):
        for name, value in list(tag.attrs.items()):
            v = " ".join(value) if isinstance(value, list) else str(value)
            low = "".join(v.split()).lower()
            if (
                name.lower() not in _SAFE_ATTRS
                or name == "href" and not low.startswith(("http:", "https:", "mailto:", "#"))
                or name == "src" and not low.startswith("data:image/")
                or name == "style" and any(k in low for k in ("url(", "expression(", "@import", "javascript:"))
            ):
                del tag.attrs[name]


def _html_table(table) -> list[TableCell]:
    cells: list[TableCell] = []
    occupied: set[tuple[int, int]] = set()
    rows = [tr for tr in table.find_all("tr") if tr.find_parent("table") is table]
    for r, tr in enumerate(rows):
        c = 0
        for cell in tr.find_all(["td", "th"], recursive=False):
            while (r, c) in occupied:
                c += 1
            rs, cs = int(cell.get("rowspan", 1) or 1), int(cell.get("colspan", 1) or 1)
            for i in range(r, r + rs):
                for k in range(c, c + cs):
                    occupied.add((i, k))
            cells.append(TableCell(row=r, col=c, rowspan=rs, colspan=cs, text=" ".join(cell.get_text(" ").split())))
            c += cs
    return cells


def _parse_html(data: bytes) -> list[tuple[Page, str]]:
    from bs4 import BeautifulSoup, NavigableString

    soup = BeautifulSoup(data, "lxml")
    for t in soup(["script", "style", "noscript", "template"]):
        t.decompose()
    _sanitize(soup)
    items: list[Item] = []
    seen_h1 = False

    def text_of(node) -> str:
        return " ".join(node.get_text(" ").split())

    def add(type_: ElementType, text: str, html: str, level: int | None = None) -> None:
        items.append((_text_el(type_, text, level), html, ""))

    def walk(node) -> None:
        nonlocal seen_h1
        for child in node.children:
            if isinstance(child, NavigableString):
                t = " ".join(str(child).split())
                if t and node.name in _BLOCK_CONTAINERS:
                    add(ElementType.PARAGRAPH, t, f"<p>{escape(t)}</p>")
                continue
            name = child.name
            if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                level = int(name[1])
                t = text_of(child)
                if not t:
                    continue
                if level == 1 and not seen_h1:
                    seen_h1 = True
                    add(ElementType.TITLE, t, str(child), 0)
                else:
                    add(ElementType.HEADING, t, str(child), max(1, level - 1))
            elif name == "table":
                el, html = _table(_html_table(child))
                if el:
                    items.append((el, html, ""))
            elif name == "li":
                if t := text_of(child):
                    add(ElementType.LIST_ITEM, t, f"<ul>{child}</ul>")
            elif name in ("p", "pre", "address", "dd", "dt", "figcaption"):
                if t := text_of(child):
                    add(ElementType.PARAGRAPH, t, str(child))
            elif name in _BLOCK_CONTAINERS or name in ("span", "a", "b", "strong", "i", "em", "font"):
                if name in _BLOCK_CONTAINERS:
                    walk(child)
                elif t := text_of(child):
                    add(ElementType.PARAGRAPH, t, str(child))

    walk(soup.body or soup)
    return [_page(1, items)]


_PARSERS = {FileKind.DOCX: _parse_docx, FileKind.XLSX: _parse_xlsx, FileKind.PPTX: _parse_pptx, FileKind.HTML: _parse_html}


def parse_office_with_previews(data: bytes, filename: str, kind: FileKind, document_id: str, job_id: str) -> tuple[Document, dict[int, str]]:
    """AST plus one standalone HTML preview of the original file per page (keyed by page number)."""
    if kind not in _PARSERS:
        raise ValueError(f"không phải tài liệu Office/HTML: {kind}")
    parsed = _PARSERS[kind](data)
    pages = [p for p, _ in parsed]
    doc = Document(
        document_id=document_id,
        job_id=job_id,
        status=DocStatus.COMPLETED,
        metadata=DocumentMetadata(filename=filename, file_type=MIME[kind], total_pages=len(pages), engine_versions={"parser": f"native-{kind.value}"}),
        pages=pages,
    )
    return doc, {p.page_number: _preview_document(body, kind) for p, body in parsed}


def parse_office(data: bytes, filename: str, kind: FileKind, document_id: str, job_id: str) -> Document:
    return parse_office_with_previews(data, filename, kind, document_id, job_id)[0]
