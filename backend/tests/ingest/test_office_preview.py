"""Original-file preview for the Office/HTML fast path (no page raster exists for these formats)."""

import io
import re

import docx
import openpyxl
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches

from idp.ingest.detect import FileKind
from idp.ingest.office import parse_office_with_previews

from .test_office import _docx_bytes


def _ids(doc) -> dict[int, list[str]]:
    return {p.page_number: [e.element_id for e in p.elements] for p in doc.pages}


def _assert_every_element_linked(doc, previews):
    assert set(previews) == {p.page_number for p in doc.pages}
    for n, ids in _ids(doc).items():
        for el_id in ids:
            assert f'data-el="{el_id}"' in previews[n], el_id


def _block(html: str, el_id: str) -> str:
    """Markup of the block tagged with ``el_id`` (up to the next tagged block)."""
    start = html.index(f'data-el="{el_id}"')
    nxt = html.find("data-el=", start + 10)
    return html[start : nxt if nxt != -1 else len(html)]


def test_docx_preview_keeps_formatting_and_links_elements():
    d = docx.Document()
    d.add_heading("BÁO CÁO", level=1)
    p = d.add_paragraph("Số tiền: ")
    p.add_run("15.000.000").bold = True
    d.add_paragraph("<script>alert(1)</script>")
    buf = io.BytesIO()
    d.save(buf)
    doc, previews = parse_office_with_previews(buf.getvalue(), "a.docx", FileKind.DOCX, "d", "j")
    _assert_every_element_linked(doc, previews)
    html = previews[1]
    assert html.lstrip().lower().startswith("<!doctype html>")
    assert "<b>15.000.000</b>" in _block(html, "el_p1_002")
    assert "<script>" not in html and "&lt;script&gt;" in html


def test_docx_preview_table_spans_and_cell_coordinates():
    doc, previews = parse_office_with_previews(_docx_bytes(), "bctc.docx", FileKind.DOCX, "d", "j")
    _assert_every_element_linked(doc, previews)
    table = next(e for e in doc.pages[0].elements if e.type.value == "table")
    block = _block(previews[1], table.element_id)
    assert re.search(r'<td[^>]*rowspan="2"[^>]*data-r="0" data-c="0"', block)
    assert re.search(r'<td[^>]*colspan="2"[^>]*data-r="0" data-c="1"', block)
    assert "13.870" in block


def test_xlsx_preview_one_page_per_sheet_with_styles():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BCTC"
    ws["A1"] = "Chỉ tiêu"
    ws["A1"].font = Font(bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B1"] = "Năm 2026"
    ws.merge_cells("B1:C1")
    ws["A2"] = "Doanh thu"
    ws["B2"] = 12450000
    ws2 = wb.create_sheet("Khác")
    ws2["A1"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    doc, previews = parse_office_with_previews(buf.getvalue(), "a.xlsx", FileKind.XLSX, "d", "j")
    _assert_every_element_linked(doc, previews)
    html = previews[1]
    assert "font-weight:bold" in html and "background:#FFFF00" in html
    assert re.search(r'colspan="2"[^>]*data-r="0" data-c="1"', html)
    assert "12.450.000" in html
    assert "Khác" not in html and "Khác" in previews[2]


def test_pptx_preview_positions_shapes_on_slide():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Kết quả"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    box.text_frame.text = "Doanh thu tăng 12%"
    buf = io.BytesIO()
    prs.save(buf)
    doc, previews = parse_office_with_previews(buf.getvalue(), "a.pptx", FileKind.PPTX, "d", "j")
    _assert_every_element_linked(doc, previews)
    html = previews[1]
    assert 'class="slide"' in html and "aspect-ratio" in html
    el_id = next(e.element_id for e in doc.pages[0].elements if e.text == "Doanh thu tăng 12%")
    assert re.search(r"left:10(\.0+)?%", _block(html, el_id))  # 1in of a 10in-wide slide


def test_html_preview_keeps_markup_but_strips_active_content():
    src = """<html><body><h1 onclick="x()">Hợp đồng</h1>
    <p>Điều khoản <strong>chung</strong> <a href="javascript:alert(1)">link</a></p>
    <iframe src="http://evil"></iframe><script>steal()</script>
    <table><tr><td colspan="2">A</td></tr></table></body></html>""".encode()
    doc, previews = parse_office_with_previews(src, "a.html", FileKind.HTML, "d", "j")
    _assert_every_element_linked(doc, previews)
    html = previews[1]
    assert "<strong>chung</strong>" in html
    assert "onclick" not in html and "javascript:" not in html
    assert "<iframe" not in html and "steal()" not in html
    assert re.search(r'colspan="2"[^>]*data-r="0" data-c="0"', html)
