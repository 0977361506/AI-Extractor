import io

import docx
import openpyxl
import pytest
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches

from idp.core.ast import ElementType
from idp.ingest.detect import FileKind
from idp.ingest.office import parse_office


def _docx_bytes() -> bytes:
    d = docx.Document()
    d.add_heading("BÁO CÁO TÀI CHÍNH", level=1)
    d.add_paragraph("Đơn vị tính: triệu đồng")
    d.add_paragraph("Mục một", style="List Bullet")
    legacy = d.add_paragraph().add_run("Ng©n hµng")
    legacy.font.name = ".VnTime"
    legacy.font.size = Pt(12)
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(1, 0)).text = "Chỉ tiêu"  # vertical merge
    t.cell(0, 1).merge(t.cell(0, 2)).text = "Năm 2026"  # horizontal merge
    t.cell(1, 1).text = "Quý I"
    t.cell(1, 2).text = "Quý II"
    t.cell(2, 0).text = "Doanh thu"
    t.cell(2, 1).text = "12.450"
    t.cell(2, 2).text = "13.870"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_headings_paragraphs_legacy_font_and_merged_table():
    doc = parse_office(_docx_bytes(), "bctc.docx", FileKind.DOCX, "d", "j")
    els = doc.pages[0].elements
    assert els[0].type == ElementType.HEADING and els[0].text == "BÁO CÁO TÀI CHÍNH"
    assert els[2].type == ElementType.LIST_ITEM
    assert any(e.text == "Ngân hàng" for e in els)
    table = next(e for e in els if e.type == ElementType.TABLE)
    spans = {(c.row, c.col): (c.rowspan, c.colspan, c.text) for c in table.cells}
    assert spans[(0, 0)] == (2, 1, "Chỉ tiêu")
    assert spans[(0, 1)] == (1, 2, "Năm 2026")
    assert spans[(2, 2)] == (1, 1, "13.870")
    assert len(table.cells) == 7
    assert [e.reading_order for e in els] == list(range(len(els)))


def test_xlsx_sheets_merged_cells_and_number_format():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BCTC"
    ws["A1"] = "Chỉ tiêu"
    ws["B1"] = "Năm 2026"
    ws.merge_cells("B1:C1")
    ws["A2"] = "Doanh thu"
    ws["B2"] = 12450000
    ws["C2"] = -1290.5
    wb.create_sheet("Trống")
    buf = io.BytesIO()
    wb.save(buf)
    doc = parse_office(buf.getvalue(), "bctc.xlsx", FileKind.XLSX, "d", "j")
    assert len(doc.pages) == 1  # empty sheet skipped
    table = next(e for e in doc.pages[0].elements if e.type == ElementType.TABLE)
    spans = {(c.row, c.col): (c.colspan, c.text) for c in table.cells}
    assert spans[(0, 1)] == (2, "Năm 2026")
    assert spans[(1, 1)][1] == "12.450.000"
    assert spans[(1, 2)][1] == "-1.290,5"
    assert doc.pages[0].elements[0].text == "BCTC"


def test_pptx_slides_title_text_and_table():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Kết quả kinh doanh"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    box.text_frame.text = "Doanh thu tăng 12%"
    tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    tbl.cell(0, 0).merge(tbl.cell(0, 1))
    tbl.cell(0, 0).text = "Quý III"
    tbl.cell(1, 0).text = "Doanh thu"
    tbl.cell(1, 1).text = "15.220"
    buf = io.BytesIO()
    prs.save(buf)
    doc = parse_office(buf.getvalue(), "kq.pptx", FileKind.PPTX, "d", "j")
    els = doc.pages[0].elements
    assert els[0].type == ElementType.TITLE and els[0].text == "Kết quả kinh doanh"
    assert any(e.text == "Doanh thu tăng 12%" for e in els)
    table = next(e for e in els if e.type == ElementType.TABLE)
    assert {(c.row, c.col): c.colspan for c in table.cells}[(0, 0)] == 2


def test_html_structure():
    html = """<html><head><style>p{}</style><script>x()</script></head><body>
    <h1>Hợp đồng</h1><p>Điều khoản <b>chung</b></p><ul><li>Mục 1</li></ul>
    <table><tr><th rowspan="2">Chỉ tiêu</th><th colspan="2">2026</th></tr><tr><th>Q1</th><th>Q2</th></tr>
    <tr><td>DT</td><td>1</td><td>2</td></tr></table></body></html>""".encode()
    doc = parse_office(html, "a.html", FileKind.HTML, "d", "j")
    els = doc.pages[0].elements
    assert [e.type for e in els][:3] == [ElementType.TITLE, ElementType.PARAGRAPH, ElementType.LIST_ITEM]
    assert els[1].text == "Điều khoản chung"
    table = next(e for e in els if e.type == ElementType.TABLE)
    assert {(c.row, c.col): (c.rowspan, c.colspan) for c in table.cells}[(0, 0)] == (2, 1)
    assert all("x()" not in e.text for e in els)


def test_unsupported_kind_raises():
    with pytest.raises(ValueError):
        parse_office(b"x", "a.bin", FileKind.UNKNOWN, "d", "j")
