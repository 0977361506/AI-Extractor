import json

import httpx
import pytest
from fastapi.testclient import TestClient

from idp.api.app import create_app
from idp.core.ast import Document, DocumentMetadata, Element, ElementType, Page, PageImages, ReviewReason, TableCell
from idp.core.config import Settings
from idp.jobs.db import Database
from idp.jobs.runner import InlineRunner
from idp.storage.local import LocalStorage


def fake_job(document_id, job_id, filename, data_dir, db_path):
    """Stand-in for the OCR pipeline: writes a small AST + progress events."""
    import numpy as np

    db = Database(db_path)
    st = LocalStorage(data_dir)
    st.put_image(document_id, "pages/1_display.jpg", np.full((20, 10, 3), 255, np.uint8))
    cells = [
        TableCell(row=0, col=0, text="Chỉ tiêu", is_header=True),
        TableCell(row=0, col=1, text="Số tiền", is_header=True),
        TableCell(row=1, col=0, text="Tiền"),
        TableCell(row=1, col=1, text="15.250.0O0", confidence=0.6),
    ]
    page = Page(
        page_number=1,
        width=10,
        height=20,
        images=PageImages(raw="x", display=f"{document_id}/pages/1_display.jpg", ocr="x"),
        elements=[
            Element(element_id="el_p1_001", type=ElementType.TITLE, reading_order=0, text="HỢP ĐỒNG", content_markdown="# HỢP ĐỒNG"),
            Element(element_id="el_p1_002", type=ElementType.TABLE, reading_order=1, cells=cells, needs_human_review=True,
                    review_reasons=[ReviewReason.LOW_CONFIDENCE]),
        ],
    )
    st.put_text(document_id, "pages/1.json", page.model_dump_json())
    db.add_event(document_id, {"document_id": document_id, "status": "IN_PROGRESS", "progress_percentage": 50, "message_vi": "trang 1"})
    doc = Document(document_id=document_id, job_id=job_id, metadata=DocumentMetadata(filename=filename, file_type="application/pdf", total_pages=1), pages=[page])
    from idp.pipeline.document import finalize_document

    finalize_document(doc, st)
    db.update_document(document_id, status="COMPLETED", done_pages=1, total_pages=1)
    db.add_event(document_id, {"document_id": document_id, "status": "COMPLETED", "progress_percentage": 100, "message_vi": "xong"})
    return "COMPLETED"


@pytest.fixture
def client(tmp_path):
    settings = Settings(data_dir=tmp_path / "data", db_path=tmp_path / "idp.sqlite", gt_dir=tmp_path / "gt")
    app = create_app(settings=settings, runner=InlineRunner(fake_job))
    with TestClient(app) as c:
        yield c


def _upload(client, name="hd.pdf", data=b"%PDF-1.7 test"):
    r = client.post("/api/v1/documents", files=[("files", (name, data, "application/pdf"))])
    assert r.status_code == 200, r.text
    return r.json()["documents"][0]


def test_upload_list_and_get(client):
    d = _upload(client)
    assert d["document_id"].startswith("doc_") and d["status"] == "PENDING"
    assert [x["document_id"] for x in client.get("/api/v1/documents").json()] == [d["document_id"]]
    doc = client.get(f"/api/v1/documents/{d['document_id']}").json()
    assert doc["status"] == "COMPLETED" and doc["ast"]["pages"][0]["elements"][0]["text"] == "HỢP ĐỒNG"


def test_unsupported_file_rejected(client):
    r = client.post("/api/v1/documents", files=[("files", ("x.exe", b"MZ....", "application/octet-stream"))])
    assert r.json()["documents"][0]["error"]


def test_page_and_image(client):
    d = _upload(client)
    page = client.get(f"/api/v1/documents/{d['document_id']}/pages/1").json()
    assert page["page_number"] == 1
    img = client.get(f"/api/v1/documents/{d['document_id']}/pages/1/image?kind=display")
    assert img.status_code == 200 and img.headers["content-type"].startswith("image/")
    assert client.get(f"/api/v1/documents/{d['document_id']}/pages/9").status_code == 404


def test_sse_replays_events_until_done(client):
    d = _upload(client)
    events = []
    with client.stream("GET", f"/api/v1/documents/{d['document_id']}/events") as r:
        for line in r.iter_lines():
            if line.startswith("data:"):
                events.append(json.loads(line[5:]))
            if line.startswith("event: end"):
                break
    assert [e.get("progress_percentage") for e in events if "progress_percentage" in e][-1] == 100


def test_corrections_audit_export_and_ground_truth(client, tmp_path):
    d = _upload(client)
    doc_id = d["document_id"]
    r = client.post(f"/api/v1/documents/{doc_id}/corrections", json={"element_id": "el_p1_002", "path": "cells[1][1].text", "value": "15.250.000"})
    assert r.status_code == 200, r.text
    assert next(c for c in r.json()["element"]["cells"] if c["row"] == 1 and c["col"] == 1)["text"] == "15.250.000"
    audit = client.get(f"/api/v1/documents/{doc_id}/corrections").json()
    assert audit[0]["old_value"] == "15.250.0O0" and audit[0]["editor"]
    client.post(f"/api/v1/documents/{doc_id}/corrections", json={"element_id": "el_p1_002", "path": "review", "value": True})
    ast = client.get(f"/api/v1/documents/{doc_id}").json()["ast"]
    table = ast["pages"][0]["elements"][1]
    assert table["human_corrected"] and not table["needs_human_review"]
    md = client.get(f"/api/v1/documents/{doc_id}/export?format=md").text
    assert "15.250.000" in md
    jsonl = client.get(f"/api/v1/documents/{doc_id}/export?format=jsonl").text
    assert json.loads(jsonl.splitlines()[0])["chunk_id"]
    gt = client.post(f"/api/v1/documents/{doc_id}/ground-truth").json()
    saved = json.loads(open(gt["path"], encoding="utf-8").read())
    assert saved["pages"][0]["elements"][1]["cells"][3]["text"] == "15.250.000"
    bad = client.post(f"/api/v1/documents/{doc_id}/corrections", json={"element_id": "el_p1_002", "path": "__class__", "value": "x"})
    assert bad.status_code == 400


def test_vlm_settings_masked_and_merge(client):
    r = client.put("/api/v1/settings/vlm", json={"enabled": True, "base_url": "https://gw/v1", "api_key": "sk-secret-9999", "model": "qwen"})
    assert r.json()["api_key"] == "sk-****9999"
    r = client.put("/api/v1/settings/vlm", json={**r.json(), "model": "qwen3-vl"})
    body = client.get("/api/v1/settings/vlm").json()
    assert body["model"] == "qwen3-vl" and body["api_key"] == "sk-****9999" and body["api_key_set"]


def test_vlm_connection_test_endpoint_uses_unsaved_settings(tmp_path):
    def gateway(request: httpx.Request):
        if request.url.path.endswith("/models"):
            return httpx.Response(401, json={"error": "bad key"})
        return httpx.Response(401, json={"error": "bad key"})

    settings = Settings(data_dir=tmp_path / "data", db_path=tmp_path / "idp.sqlite", gt_dir=tmp_path / "gt")
    app = create_app(settings=settings, runner=InlineRunner(fake_job), vlm_transport=httpx.MockTransport(gateway))
    with TestClient(app) as c:
        r = c.post("/api/v1/settings/vlm/test", json={"base_url": "https://gw/v1", "api_key": "wrong", "model": "m"})
        result = r.json()
        assert not result["ok"]
        assert any(s["name"] == "models" and "401" in s["detail"] for s in result["steps"])
        assert c.get("/api/v1/settings/vlm/tests").json()[0]["result"]["ok"] is False
        assert c.get("/api/v1/settings/vlm").json()["base_url"] == ""  # test does not save


def test_delete(client):
    d = _upload(client)
    assert client.delete(f"/api/v1/documents/{d['document_id']}").status_code == 200
    assert client.get(f"/api/v1/documents/{d['document_id']}").status_code == 404


@pytest.fixture
def office_client(tmp_path):
    settings = Settings(data_dir=tmp_path / "data", db_path=tmp_path / "idp.sqlite", gt_dir=tmp_path / "gt")
    with TestClient(create_app(settings=settings, runner=InlineRunner())) as c:
        yield c, tmp_path / "data"


def _docx() -> bytes:
    import io

    import docx

    d = docx.Document()
    d.add_heading("HỢP ĐỒNG TÍN DỤNG", level=1)
    d.add_paragraph("Bên vay cam kết")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_office_document_serves_original_preview_instead_of_image(office_client):
    client, data_dir = office_client
    d = _upload(client, "hd.docx", _docx())
    doc = client.get(f"/api/v1/documents/{d['document_id']}").json()
    assert doc["status"] == "COMPLETED"
    first = doc["ast"]["pages"][0]["elements"][0]["element_id"]
    r = client.get(f"/api/v1/documents/{d['document_id']}/pages/1/preview")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    assert "default-src 'none'" in r.headers["content-security-policy"]
    assert f'data-el="{first}"' in r.text and "HỢP ĐỒNG TÍN DỤNG" in r.text
    assert client.get(f"/api/v1/documents/{d['document_id']}/pages/2/preview").status_code == 404


def test_preview_backfilled_for_documents_processed_before_previews_existed(office_client):
    client, data_dir = office_client
    d = _upload(client, "hd.docx", _docx())
    stored = data_dir / d["document_id"] / "pages" / "1_preview.html"
    assert stored.exists()
    stored.unlink()
    r = client.get(f"/api/v1/documents/{d['document_id']}/pages/1/preview")
    assert r.status_code == 200 and "Bên vay cam kết" in r.text
    assert stored.exists()


def test_scanned_document_has_no_preview(client):
    d = _upload(client)
    assert client.get(f"/api/v1/documents/{d['document_id']}/pages/1/preview").status_code == 404
